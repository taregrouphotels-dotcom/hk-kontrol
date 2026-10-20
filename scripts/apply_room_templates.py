#!/usr/bin/env python3
# Elle hazırlanan oda tipi şablonlarını (templates/X1.json gibi) data/HK_Veri_Toplama.xlsx'e yazar:
# Kontrol Adımları, Envanter - Demirbaş / Mutfak, Minibar ve İkram ve eksik Kalem Kataloğu satırları.
# Şablondaki oda tiplerinin eski satırları silinip yeniden yazılır; diğer satırlara dokunulmaz.
# Kullanım: python3 scripts/apply_room_templates.py data/HK_Veri_Toplama.xlsx templates/X1.json templates/X2.json
import json, re, sys, zipfile, shutil, tempfile
from copy import copy
import openpyxl

path, tpl_paths = sys.argv[1], sys.argv[2:]
tpls = [json.load(open(p)) for p in tpl_paths]
codes = {t['code'] for t in tpls}
original = tempfile.mktemp(suffix='.xlsx')
shutil.copy(path, original)
wb = openpyxl.load_workbook(path)
Y, N = 'Evet', 'Hayır'


def header(ws):
    return {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}


def delete_rows(ws, pred):
    for r in range(ws.max_row, 1, -1):
        if pred([ws.cell(r, c).value for c in range(1, ws.max_column + 1)]):
            ws.delete_rows(r)


def last_row(ws):
    return max([r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value] or [1])


def append(ws, values, style_row=2):
    r = last_row(ws) + 1
    for c, v in values.items():
        cell = ws.cell(r, c, v)
        cell._style = copy(ws.cell(style_row, c)._style)
    return r


# ---- Kalem kataloğu: eksik kalemler ----
ks = wb['Kalem Kataloğu']
kh = header(ks)
have = {(r[0], r[1]) for r in ks.iter_rows(min_row=2, values_only=True)}
kinds = {(r[0], r[1]): r[3] for r in ks.iter_rows(min_row=2, values_only=True)}
added = 0
for t in tpls:
    for a in t.get('catalog_additions', []):
        if (t['property'], a['key']) in have: continue
        append(ks, {kh['Tesis']: t['property'], kh['Kalem Kodu']: a['key'], kh['Kalem Adı']: a['label'], kh['Tür']: a['kind'],
                    kh['Tarif (ayırt edici)']: a.get('description'), kh['Kaynak']: 'Drive foto (20.10.2026)'})
        have.add((t['property'], a['key'])); kinds[(t['property'], a['key'])] = a['kind']; added += 1

# ---- Kontrol adımları ----
ss = wb['Kontrol Adımları']
sh = header(ss)
delete_rows(ss, lambda v: v[0] in codes)
for t in tpls:
    for s in t['steps']:
        append(ss, {sh['Oda Tipi Kodu']: t['code'], sh['Sıra']: s['order'], sh['Bölge']: s['zone'], sh['Adım Adı']: s['title'],
                    sh['Çekim Talimatı']: s['instruction'], sh['Kontrol Maddeleri (her satır bir madde)']: '\n'.join(s['checks']),
                    sh['Çıkışta']: Y if 'full' in s['applies_to'] else N, sh['Konaklamada']: Y if 'stayover' in s['applies_to'] else N,
                    sh['Fener Gerekli']: Y if s['needs_flash'] else N, sh['Zorunlu']: Y if s['required'] else N,
                    sh['Kare Sayısı']: s['shots'], sh['AI Talimatı']: s['ai_prompt']})

# ---- Envanter ----
F_STEP = "=IFERROR(INDEX('Kontrol Adımları'!$D:$D,MATCH(1,INDEX(('Kontrol Adımları'!$A$1:$A$400=B{r})*('Kontrol Adımları'!$B$1:$B$400=C{r}),0),0)),\"\")"
F_CAT = "=IFERROR(INDEX('Kalem Kataloğu'!${col}$1:${col}$1000,MATCH(1,INDEX(('Kalem Kataloğu'!$A$1:$A$1000=A{r})*('Kalem Kataloğu'!$B$1:$B$1000=E{r}),0),0)),\"?\")"
inv = {name: wb[name] for name in ('Envanter - Demirbaş', 'Envanter - Mutfak')}
for ws in inv.values():
    delete_rows(ws, lambda v: v[1] in codes and v[0] in {t['property'] for t in tpls})
treats = wb['Minibar ve İkram']
th = header(treats)
delete_rows(treats, lambda v: v[1] in codes and v[0] in {t['property'] for t in tpls})
n_inv = n_treat = 0
for t in tpls:
    seen = set()
    for s in t['steps']:
        for i in s['items']:
            if i['key'] in seen: raise SystemExit(f"{t['code']}: {i['key']} iki kez")
            seen.add(i['key'])
            kind = kinds.get((t['property'], i['key']))
            if kind is None: raise SystemExit(f"{t['code']}: {i['key']} katalogda yok")
            if kind in ('İkram', 'İçecek'):
                append(treats, {th['Tesis']: t['property'], th['Oda Tipi Kodu']: t['code'], th['Adım No']: s['order'], th['Yer']: s['title'],
                                th['Kalem Kodu']: i['key'], th['Ürün']: i['label'], th['Standart Adet']: i['qty'], th['Not']: i.get('note')})
                n_treat += 1; continue
            ws = inv['Envanter - Mutfak' if kind == 'Mutfak Malzemesi' else 'Envanter - Demirbaş']
            h = header(ws)
            append(ws, {h['Tesis']: t['property'], h['Oda Tipi Kodu']: t['code'], h['Adım No']: s['order'], h['Kalem Kodu']: i['key'],
                        h['Olması Gereken Adet']: i['qty'], h['Doğrulandı mı']: Y if i['verified'] else N,
                        h['Kaynak']: 'Drive foto (20.10.2026)', h['Not']: i.get('note')})
            n_inv += 1
    ikram_step = next((s['order'] for s in t['steps'] if 'ikram' in s['title'].lower()), None)
    for x in t.get('treats', []):
        append(treats, {th['Tesis']: t['property'], th['Oda Tipi Kodu']: t['code'], th['Adım No']: ikram_step, th['Yer']: 'İkram tepsisi',
                        th['Kalem Kodu']: x['key'], th['Ürün']: x['label'], th['Standart Adet']: x['qty'], th['Not']: t.get('treats_note')})
        n_treat += 1
# Satır silinince kayan formülleri her satır için yeniden yaz
for ws in inv.values():
    h = header(ws)
    for r in range(2, last_row(ws) + 1):
        ws.cell(r, h['Adım Adı (otomatik)'], F_STEP.format(r=r))
        ws.cell(r, h['Kalem Adı (otomatik)'], F_CAT.format(col='C', r=r))
        ws.cell(r, h['Tür (otomatik)'], F_CAT.format(col='D', r=r))
wb.save(path)


# openpyxl sayfa düzeyindeki açılır liste uzantısını (extLst) siler: orijinal dosyadan geri koy
def sheet_files(z):
    wbx = z.read('xl/workbook.xml').decode()
    rels = z.read('xl/_rels/workbook.xml.rels').decode()
    tgt = dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    tgt.update({a: b for b, a in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        t = tgt[rid].lstrip('/')
        out[name] = t if t.startswith('xl/') else 'xl/' + t
    return out


zo = zipfile.ZipFile(original)
zn = zipfile.ZipFile(path)
fo, fn = sheet_files(zo), sheet_files(zn)
files = {i.filename: zn.read(i.filename) for i in zn.infolist()}
restored = 0
for name, f in fo.items():
    m = re.search(r'<extLst>.*</extLst>', zo.read(f).decode(), re.S)
    if not m or name not in fn: continue
    x = files[fn[name]].decode()
    if '<extLst>' in x: continue
    files[fn[name]] = x.replace('</worksheet>', m.group(0) + '</worksheet>').encode()
    restored += 1
zn.close()
with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
    for k, v in files.items(): z.writestr(k, v)
print(f'katalog +{added}, adım {sum(len(t["steps"]) for t in tpls)}, envanter {n_inv}, ikram {n_treat}, açılır liste geri konan sayfa {restored}')
